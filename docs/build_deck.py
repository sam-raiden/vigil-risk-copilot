"""Builds docs/Vigil_deck.pptx. Run: python docs/build_deck.py (needs python-pptx)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
TEAL = RGBColor(0x00, 0x9E, 0xA8)
INK = RGBColor(0x1F, 0x29, 0x37)
GREY = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def box(slide, x, y, w, h, fill=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.line.fill.background()
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    return s


def text(slide, x, y, w, h, lines, size=18, color=INK, bold=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = bold
        p.space_after = Pt(8)
    return tb


def slide(title, bullets, note=None):
    s = prs.slides.add_slide(BLANK)
    box(s, 0, 0, 13.333, 1.1, NAVY)
    text(s, 0.6, 0.25, 12, 0.7, [title], size=30, color=WHITE, bold=True)
    text(s, 0.6, 1.5, 12.1, 5, bullets, size=20)
    if note:
        text(s, 0.6, 6.7, 12.1, 0.6, [note], size=13, color=GREY)
    return s


# 1 title
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, 13.333, 7.5, NAVY)
box(s, 0.6, 3.55, 1.6, 0.06, TEAL)
text(s, 0.6, 2.0, 12, 1.4, ["Vigil"], size=60, color=WHITE, bold=True)
text(s, 0.6, 3.8, 12, 1.5, ["A cited risk, fraud and regulatory copilot for banking and NBFC compliance",
                            "Built with Snowflake CoCo | synthetic data | GCC Edition hackathon"], size=22, color=WHITE)

slide("The problem", [
    "Compliance teams get an alert, then assemble the evidence and the matching policy clause by hand.",
    "Fraud rings are harder: each account looks ordinary on its own, so per-transaction rules never fire.",
    "Goal: signal -> evidence -> documented finding, from a plain-English question.",
])

slide("What Vigil does", [
    "Fraud / AML: structuring, high-risk remittances, velocity, and ring detection over a relationship graph.",
    "Liquidity: LCR calculation and breach history with the Basel III clause.",
    "Credit: NPA classification and provisioning with the RBI IRAC clause.",
    "Every answer cites a record ID and a policy document + clause, or says it cannot.",
    "A copilot: a human signs off on every filing; nothing acts unasked.",
])

slide("Architecture (all in Snowflake)", [
    "8 tables / Dynamic Tables in VIGIL.CORE: customers, accounts, transactions, alerts, LCR snapshots, credit profiles, RING_EDGES, RINGS.",
    "Cortex Search over 35 numbered policy clauses; semantic view with 6 verified queries.",
    "One Cortex Agent, three tools: analytics, policy search, ring lookup.",
    "Streamlit-in-Snowflake app: chat + audit panel + ring graph + markdown download.",
])

slide("The novel piece: ring detection", [
    "RING_EDGES: pairs of different accounts that paid the same counterparty within 72 hours.",
    "Hub filter: drop counterparties paid by more than 8 distinct accounts (utilities).",
    "RINGS: connected components of 3+ accounts (5 min-label-propagation steps; Dynamic Tables cannot recurse).",
    "Measured on the seeded data: 2,436 edges without the filter, 10 with it, exactly one ring: ACC-0031..0035 via 'Orion Trading Co'.",
    "The seeded utility (25 payers) is not reported as a ring.",
], note="Threshold 8 is tuned on one synthetic case.")

slide("Policy-grounded answers", [
    "Six policy documents, 35 numbered clauses, including POL-AML-001 Clause 4.5 (RING_PATTERN).",
    "Agent rules: cite record ID + clause; state uncertainty; ask which account/loan when unspecified; always check the ring lookup for fraud questions.",
    "No external tool exists, so the agent cannot post anywhere; it says a human must.",
    "The app shows a 'Not citable' warning instead of presenting an uncited answer as a finding.",
])

slide("Built through CoCo, tested for real", [
    "Every Snowflake object was generated and run through CoCo; prompts and results are logged per phase in the repo.",
    "Semantic view: 6 verified questions, 0 mismatches against seeded facts (CoCo's check).",
    "Agent: 10 scripted calls, 9 clean passes + 1 with a caveat (LN-0033 clause cited only via search annotations).",
    "Caught by opening the app: a first-build UI crash that CoCo's backend-only tests had missed; fixed through CoCo.",
])

slide("Measured timings", [
    "Single-turn agent calls, wall-clock around the call (n = 1 each):",
    "Ambiguous question 5.7 s | NPA 21.4 s | LCR 29.2 s | structuring 31.0 s | velocity 31.5 s",
    "Ring 32.6 s | high-risk remittance 32.8 s | utility check 33.3 s",
    "The app itself measured 28.4 s for the ring question from the UI.",
    "No manual baseline was measured, so no speed-up claim is made.",
], note="Single runs, not averages; model latency varies.")

slide("Honest limits", [
    "Small synthetic dataset (hundreds of transactions); threshold tuned on one case.",
    "Exact-name matching only; no device attribute; no EDD record table (EDD shown as unknown).",
    "No live regulatory feed; no second agent or external connector; nothing posts externally.",
    "The app runs inside Snowflake and needs a login. A human signs off on every filing.",
])

slide("Next steps", [
    "Fuzzy entity resolution and a device/IP attribute for the graph.",
    "Calibrate the hub threshold on a larger synthetic set.",
    "Add an EDD record table so EDD status is a fact, not 'unknown'.",
    "Scale and cost testing before any real-data pilot.",
])

prs.save("docs/Vigil_deck.pptx")
print("saved", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
